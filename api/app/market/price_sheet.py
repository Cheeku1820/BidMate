"""The price request an estimator sends a supplier, and the parser
that reads the filled sheet back (estimate-first-pricing §6). Pure:
bytes in, records out. Only two cells are interpreted -- the price (a
number) and the row key (an id the caller checks belongs to the
project). Everything else is text, carried as text.

A row the parser can't read is reported, never raised: it lands in
`ParsedSheet.unreadable` with its line and a reason in the estimator's
words, so a supplier's "call for price" or a NaN cell names its own row
in the preview instead of failing the whole upload. A blank price is
not unreadable -- it is an unpriced row, as before."""
from __future__ import annotations

import csv
import io
import re
from decimal import Decimal, InvalidOperation
from typing import NamedTuple

import openpyxl
from openpyxl.utils import get_column_letter

HEADER = ("Item", "Description", "Qty", "Unit", "Unit price", "Supplier part no.", "Notes", "Row key")
REQUIRED_HEADER = ("Item", "Unit price")
_KEY_COL = HEADER.index("Row key") + 1
# A leading or trailing currency word ("USD 9.10", "9.10 USD"), stripped
# before the number itself is validated.
_CURRENCY_WORD_LEAD_RE = re.compile(r"^[A-Za-z]+\s+")
_CURRENCY_WORD_TRAIL_RE = re.compile(r"\s+[A-Za-z]+$")
# US-style thousands grouping ("2,250.00") -- the comma is only ever
# removed once the grouping itself is confirmed correct, never blindly,
# so a European "2.250,00" (which is a different number, not a
# formatting variant) is refused rather than silently mis-parsed.
_THOUSANDS_RE = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$")
_PLAIN_NUMBER_RE = re.compile(r"^-?\d+(\.\d+)?$")
# A unit price is zero or more and below this: the column it lands in
# (ProjectMaterialPrice.price_override, Numeric(10, 2)) holds up to
# 99,999,999.99, and a negative price is a typo that would subtract
# from the bid. Either reads back as unpriced, never as a number.
PRICE_LIMIT = Decimal("100000000")

REFUSED_ALL_UNREADABLE = "None of the rows could be read. Start from Download price request."
# The reasons a row can be unreadable, as the preview lists them
# ("Row 14 -- the price isn't a number").
NOT_A_NUMBER = "the price isn't a number"
NO_ITEM_NAME = "the row has no item name"
ROW_UNREADABLE = "the row couldn't be read"


def price_in_range(value: Decimal) -> bool:
    # NaN survives quantize and raises on comparison; ask first.
    return value.is_finite() and Decimal(0) <= value < PRICE_LIMIT


class RowUnreadable(Exception):
    """A row the parser can't read; str(exc) is the reason, in the
    estimator's words."""


class ParsedRow(NamedTuple):
    row_key: str | None
    item_name: str
    unit_price: Decimal | None
    part_no: str
    notes: str
    line: int


class ParsedSheet(NamedTuple):
    rows: list[ParsedRow]
    refused: str | None
    unreadable: list[tuple[int, str]] = []   # (line, reason)


def build_request_workbook(rows: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Price request"
    ws.append(list(HEADER))
    for r in rows:
        ws.append([r["item_name"], (r.get("description") or "").splitlines()[0] if r.get("description") else "",
                   r["quantity"], r["unit"], None, None, None, str(r["item_id"])])
    ws.column_dimensions[get_column_letter(_KEY_COL)].hidden = True
    for col, width in zip("ABCDEFG", (36, 48, 8, 8, 14, 20, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _read_price(v) -> Decimal | None:
    """The one cell besides the row key that's interpreted rather than
    carried as text -- so a wrong number here is the failure this
    module exists to prevent. Blank (None or "") is None: an unpriced
    row. A number (int/float/Decimal, but not bool: `True`/`False` are
    never a price) converts directly. A string is accepted only when,
    after stripping a `$`, surrounding whitespace, and a leading or
    trailing currency word, what remains is unambiguously a plain or
    US-grouped decimal number. Anything else present in the cell --
    scientific notation ("1e3"), a European "2.250,00", "call for
    price", NaN, infinity -- raises RowUnreadable(NOT_A_NUMBER) rather
    than guessing. A number outside `price_in_range` (negative, or one
    the price column can't hold) returns None: the row is listed as
    unpriced, not carried to the apply step to fail there."""
    if isinstance(v, bool):
        return None
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        try:
            d = Decimal(str(v)).quantize(Decimal("0.01"))
        except InvalidOperation:   # inf, sNaN
            raise RowUnreadable(NOT_A_NUMBER) from None
        if not d.is_finite():      # a quiet NaN survives quantize
            raise RowUnreadable(NOT_A_NUMBER)
        return d if price_in_range(d) else None
    s = str(v).strip()
    if not s:
        return None
    s = _CURRENCY_WORD_LEAD_RE.sub("", s)
    s = _CURRENCY_WORD_TRAIL_RE.sub("", s)
    s = s.replace("$", "").strip()
    if _THOUSANDS_RE.match(s):
        s = s.replace(",", "")
    elif not _PLAIN_NUMBER_RE.match(s):
        raise RowUnreadable(NOT_A_NUMBER)
    try:
        d = Decimal(s).quantize(Decimal("0.01"))
    except InvalidOperation:
        raise RowUnreadable(NOT_A_NUMBER) from None
    return d if price_in_range(d) else None


def _price(v) -> Decimal | None:
    """`_read_price` with the refusal folded into None -- the price, or
    nothing usable."""
    try:
        return _read_price(v)
    except RowUnreadable:
        return None


def _grid(data: bytes, filename: str) -> list[list]:
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _find_header(grid: list[list]) -> tuple[int, dict[str, int]] | None:
    for i, row in enumerate(grid[:20]):
        names = {str(c).strip().lower(): j for j, c in enumerate(row) if c is not None and str(c).strip()}
        if all(h.lower() in names for h in REQUIRED_HEADER):
            return i, names
    return None


def parse_price_sheet(data: bytes, filename: str) -> ParsedSheet:
    try:
        grid = _grid(data, filename)
    except Exception:
        return ParsedSheet([], "This file couldn't be read as a spreadsheet. Upload the .xlsx or .csv the price request was sent as.")
    found = _find_header(grid)
    if found is None:
        return ParsedSheet([], f"The sheet needs a header row with the columns {' and '.join(REQUIRED_HEADER)}. "
                               f"The price request download has them in place.")
    hi, cols = found
    rows: list[ParsedRow] = []
    unreadable: list[tuple[int, str]] = []
    for n, row in enumerate(grid[hi + 1:], start=hi + 2):
        def cell(name):
            j = cols.get(name.lower())
            return row[j] if j is not None and j < len(row) else None
        if all(c is None or not str(c).strip() for c in row):
            continue   # a formatted-but-empty row, not an unreadable one
        try:
            name = cell("Item")
            if name is None or not str(name).strip():
                raise RowUnreadable(NO_ITEM_NAME)
            key = cell("Row key")
            rows.append(ParsedRow(
                row_key=str(key).strip() if key not in (None, "") else None,
                item_name=str(name).strip(),
                unit_price=_read_price(cell("Unit price")),
                part_no=str(cell("Supplier part no.") or "").strip()[:100],
                notes=str(cell("Notes") or "").strip()[:500],
                line=n,
            ))
        except RowUnreadable as exc:
            unreadable.append((n, str(exc)))
        except Exception:   # noqa: BLE001 -- one bad row is listed, never the whole upload lost
            unreadable.append((n, ROW_UNREADABLE))
    refused = REFUSED_ALL_UNREADABLE if unreadable and not rows else None
    return ParsedSheet(rows, refused, unreadable)
