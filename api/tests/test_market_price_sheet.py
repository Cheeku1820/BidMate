"""The price request workbook and the parser that reads it back --
pure, no database (estimate-first-pricing §6)."""
import io
from decimal import Decimal

import openpyxl

import pytest

from app.market.price_sheet import (HEADER, REFUSED_ALL_UNREADABLE, ParsedSheet, RowUnreadable, _price, _read_price,
                                    build_request_workbook, parse_price_sheet, price_in_range)

ROWS = [
    {"item_id": "11111111-1111-1111-1111-111111111111", "item_name": "20A duplex receptacle", "description": "Duplex Receptacle", "quantity": 14, "unit": "EA"},
    {"item_id": "22222222-2222-2222-2222-222222222222", "item_name": "Panelboard", "description": "Panel Board", "quantity": 1, "unit": "EA"},
]


def test_request_workbook_has_one_row_per_item_and_a_hidden_key_column():
    wb = openpyxl.load_workbook(io.BytesIO(build_request_workbook(ROWS)))
    ws = wb["Price request"]
    assert [c.value for c in ws[1]] == list(HEADER)
    assert ws.cell(2, 1).value == "20A duplex receptacle" and ws.cell(2, 3).value == 14
    assert ws.cell(2, 8).value == ROWS[0]["item_id"] and ws.column_dimensions["H"].hidden is True
    assert ws.cell(2, 5).value is None    # unit price left blank for the supplier


def test_parse_round_trip_reads_prices_by_row_key():
    wb = openpyxl.load_workbook(io.BytesIO(build_request_workbook(ROWS)))
    ws = wb.active
    ws.cell(2, 5).value = 9.10
    ws.cell(2, 6).value = "HBL5362"
    ws.cell(3, 5).value = "$ 2,250.00"
    out = io.BytesIO(); wb.save(out)
    parsed = parse_price_sheet(out.getvalue(), "codale.xlsx")
    assert parsed.refused is None
    assert parsed.rows[0].row_key == ROWS[0]["item_id"] and parsed.rows[0].unit_price == Decimal("9.10") and parsed.rows[0].part_no == "HBL5362"
    assert parsed.rows[1].unit_price == Decimal("2250.00")


def test_parse_blank_price_is_none_and_kept():
    wb = openpyxl.load_workbook(io.BytesIO(build_request_workbook(ROWS)))
    out = io.BytesIO(); wb.save(out)
    parsed = parse_price_sheet(out.getvalue(), "x.xlsx")
    assert [r.unit_price for r in parsed.rows] == [None, None]


def test_parse_csv_without_key_column_matches_by_name_later():
    csv = "Item,Unit price\n20A duplex receptacle,9.10\nSomething new,4\n"
    parsed = parse_price_sheet(csv.encode(), "quote.csv")
    assert parsed.refused is None
    assert parsed.rows[0].row_key is None and parsed.rows[0].item_name == "20A duplex receptacle"
    assert parsed.rows[1].unit_price == Decimal("4")


def test_parse_refuses_a_sheet_without_a_unit_price_header():
    csv = "Product,Price\nx,1\n"
    parsed = parse_price_sheet(csv.encode(), "quote.csv")
    assert parsed.rows == [] and "Unit price" in parsed.refused and "Item" in parsed.refused


def test_parse_finds_the_header_below_a_title_row():
    csv = "Codale quote 9/18\n\nItem,Unit price\n20A duplex receptacle,9.10\n"
    assert parse_price_sheet(csv.encode(), "q.csv").rows[0].unit_price == Decimal("9.10")


def test_price_refuses_scientific_notation_and_other_ambiguous_strings():
    """A wrong unit price with no refusal is the one failure this
    module exists to prevent -- scientific notation and a European
    thousands format must never be silently mis-parsed into a number."""
    assert _price("1e3") is None
    assert _price("$ 2,250.00") == Decimal("2250.00")
    assert _price("USD 9.10") == Decimal("9.10")
    assert _price("2.250,00") is None
    assert _price(True) is None
    assert _price("") is None


def test_price_refuses_a_negative_or_oversized_number_as_unpriced():
    """A negative price would subtract from the bid, and one at or over
    100,000,000 can't be stored in the price column (Numeric(10, 2)) --
    both read back as unpriced, so the preview lists the row rather
    than carrying a number the apply step would have to refuse."""
    assert _price("-9.10") is None
    assert _price(-1) is None
    assert _price("123456789012.00") is None
    assert _price(100_000_000) is None
    assert _price("99,999,999.99") == Decimal("99999999.99")
    assert _price(0) == Decimal("0.00")
    assert _price(float("inf")) is None


def test_price_in_range_is_false_for_nan_and_infinity():
    """A NaN cell used to reach this comparison and raise -- which took
    the whole price_sheet job down with the generic read copy."""
    assert price_in_range(Decimal("NaN")) is False
    assert price_in_range(Decimal("Infinity")) is False


def _csv(rows: str) -> ParsedSheet:
    return parse_price_sheet(("Item,Unit price\n" + rows).encode(), "q.csv")


def test_parse_lists_a_row_whose_price_is_not_a_number_as_unreadable():
    p = _csv("20A duplex receptacle,call for price\nPanelboard,9.10\n")
    assert [r.item_name for r in p.rows] == ["Panelboard"]
    assert p.unreadable == [(2, "the price isn't a number")] and p.refused is None


def test_a_nan_or_infinite_price_cell_is_not_a_number_rather_than_an_exception():
    """openpyxl writes a NaN as an empty cell and refuses a literal one
    at load (which lands as "couldn't be read as a spreadsheet"), so the
    value is checked where it would arrive -- the cell reader -- rather
    than through a workbook. `_price` folds the refusal into None."""
    for v in (float("nan"), float("inf"), Decimal("NaN"), Decimal("sNaN"), "nan", "1e3"):
        with pytest.raises(RowUnreadable, match="the price isn't a number"):
            _read_price(v)
        assert _price(v) is None


def test_parse_lists_a_row_with_a_price_but_no_item_name_as_unreadable_and_skips_empty_rows():
    p = _csv(",9.10\n,\nPanelboard,9.10\n")
    assert [r.item_name for r in p.rows] == ["Panelboard"]
    assert p.unreadable == [(2, "the row has no item name")]


def test_parse_reports_any_other_row_failure_rather_than_raising(monkeypatch):
    def boom(v):
        raise RuntimeError("openpyxl did something odd")
    monkeypatch.setattr("app.market.price_sheet._read_price", boom)
    p = _csv("Panelboard,9.10\n")
    assert p.rows == [] and p.unreadable == [(2, "the row couldn't be read")]
    assert p.refused == REFUSED_ALL_UNREADABLE


def test_parse_refuses_a_sheet_whose_rows_are_all_unreadable():
    p = _csv("A,call for price\nB,TBD\n")
    assert p.rows == [] and len(p.unreadable) == 2 and p.refused == REFUSED_ALL_UNREADABLE


def test_parse_blank_price_stays_unpriced_not_unreadable():
    p = _csv("Panelboard,\n")
    assert p.rows[0].unit_price is None and p.unreadable == [] and p.refused is None
