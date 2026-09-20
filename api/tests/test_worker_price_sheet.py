"""The price_sheet job turns an uploaded sheet into a preview on the
job payload -- matched by row key, unmatched listed, unpriced listed,
a foreign row key ignored (estimate-first-pricing §6)."""
import io
import uuid

import openpyxl
from sqlalchemy import select

from app.jobs import queue
from app.market.price_sheet import build_request_workbook
from app.takeoff.models import Document, Item, Job, ReviewStatus
from tests.test_worker_read import _run_all, inline  # noqa: F401


def _sheet_doc(db, project, dana, store, data, name="codale.xlsx"):
    d = Document(project_id=project.id, filename=name, doc_type="Pricing",
                 content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                 size_bytes=len(data), sha256=uuid.uuid4().hex * 2, storage_key=f"k/{uuid.uuid4()}", uploaded_by=dana.id)
    db.add(d); db.flush()
    store.put(d.storage_key, io.BytesIO(data), d.content_type, len(data))
    return d


def test_preview_matches_by_key_and_lists_the_rest(db, project, sheet, item, dana, inline, monkeypatch):
    monkeypatch.setattr("app.db.SessionLocal", lambda: db)
    other = Item(project_id=project.id, sheet_id=sheet.id, symbol="panel", name="Panelboard", system="Distribution",
                 category="Equipment", quantity=1, unit="EA", status=ReviewStatus.READY, x=1, y=1)
    db.add(other); db.flush()
    wb = openpyxl.load_workbook(io.BytesIO(build_request_workbook([
        {"item_id": str(item.id), "item_name": item.name, "description": "", "quantity": 14, "unit": "EA"},
        {"item_id": str(other.id), "item_name": other.name, "description": "", "quantity": 1, "unit": "EA"},
    ])))
    ws = wb.active
    ws.cell(2, 5).value = 9.10                       # item priced
    ws.append(["Something extra", "", 1, "EA", 4, "", "", str(uuid.uuid4())])   # foreign key, not ours
    ws.append(["Renamed thing", "", 1, "EA", 5, "", "", None])                  # no key, no name match
    out = io.BytesIO(); wb.save(out)
    d = _sheet_doc(db, project, dana, inline, out.getvalue())
    job = queue.enqueue_price_sheet(db, d, dana.id); _run_all(db)
    db.refresh(job)
    p = job.payload["preview"]
    assert job.status == "done" and p["refused"] is None
    assert [m["item_id"] for m in p["matched"]] == [str(item.id)] and p["matched"][0]["new_unit_price"] == "9.10"
    assert [u["item_name"] for u in p["unmatched"]] == ["Something extra", "Renamed thing"]
    assert [u["item_id"] for u in p["unpriced"]] == [str(other.id)]
    assert p["supplier_name"] == "codale"


def test_preview_matches_by_exact_name_without_a_key(db, project, sheet, item, dana, inline, monkeypatch):
    monkeypatch.setattr("app.db.SessionLocal", lambda: db)
    d = _sheet_doc(db, project, dana, inline, f"Item,Unit price\n{item.name},9.10\n".encode(), name="q.csv")
    job = queue.enqueue_price_sheet(db, d, dana.id); _run_all(db)
    db.refresh(job)
    assert [m["item_id"] for m in job.payload["preview"]["matched"]] == [str(item.id)]


def test_refused_sheet_completes_with_the_reason(db, project, dana, inline, monkeypatch):
    monkeypatch.setattr("app.db.SessionLocal", lambda: db)
    d = _sheet_doc(db, project, dana, inline, b"Product,Price\nx,1\n", name="q.csv")
    job = queue.enqueue_price_sheet(db, d, dana.id); _run_all(db)
    db.refresh(job)
    assert job.status == "done" and "Unit price" in job.payload["preview"]["refused"]
